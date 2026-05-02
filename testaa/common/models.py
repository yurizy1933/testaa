from django.db import models


class Project(models.Model):
    """项目模型"""
    id = models.AutoField(primary_key=True)
    project_name = models.CharField(max_length=200, verbose_name='项目名称', default='')
    description = models.TextField(blank=True, null=True, verbose_name='项目描述')
    owner = models.CharField(max_length=100, default='', verbose_name='项目负责人')
    create_time = models.DateTimeField(auto_now_add=True, verbose_name='创建时间')
    update_time = models.DateTimeField(auto_now=True, verbose_name='更新时间')
    class Meta:
        verbose_name = '项目'
        verbose_name_plural = '项目'
        ordering = ['-create_time']

    def __str__(self):
        return self.project_name


class CommonDoc(models.Model):
    """统一文档模型"""
    # 业务类型维度
    DOC_TYPE_CHOICES = [
        ('prd', 'PRD文档'),
        ('api', 'API文档'),
    ]

    # 文件格式维度（仅做存储）
    FILE_TYPE_CHOICES = [
        ('doc', 'Word文档'),
        ('docx', 'Word文档'),
        ('xlsx', 'Excel文档'),
        ('xls', 'Excel文档'),
        ('html', 'HTML文档'),
        ('htm', 'HTML文档'),
        ('pdf', 'PDF文档'),
        ('md', 'Markdown文档'),
        ('txt', '文本文档'),
    ]

    id = models.AutoField(primary_key=True)
    doc_type = models.CharField(
        max_length=10,
        choices=DOC_TYPE_CHOICES,
        verbose_name='业务类型'
    )
    file_type = models.CharField(
        max_length=10,
        choices=FILE_TYPE_CHOICES,
        verbose_name='文件格式'
    )
    filename = models.CharField(max_length=500, verbose_name='文件名')
    version = models.CharField(
        max_length=50,
        default='1.0.0',
        verbose_name='版本号'
    )
    file_content = models.TextField(blank=True, null=True, verbose_name='文件内容')
    file_path = models.FileField(upload_to='docs/', blank=True, null=True, verbose_name='文件路径')
    is_processed = models.BooleanField(
        default=False,
        verbose_name='是否已处理'
    )
    create_time = models.DateTimeField(auto_now_add=True, verbose_name='创建时间')
    update_time = models.DateTimeField(auto_now=True, verbose_name='更新时间')
    project = models.ForeignKey(
        Project,
        on_delete=models.CASCADE,
        related_name='docs',
        verbose_name='所属项目'
    )

    class Meta:
        verbose_name = '文档'
        verbose_name_plural = '文档'
        ordering = ['-create_time']
        indexes = [
            models.Index(fields=['doc_type']),
            models.Index(fields=['project', 'doc_type']),
        ]

    def __str__(self):
        return f"{self.filename} - v{self.version} - {self.get_doc_type_display()}"

    @property
    def doc_content(self):
        """获取文档内容（根据文件类型解析）"""
        if self.file_content:
            return self.file_content
        elif self.file_path:
            try:
                if self.file_type in ['doc', 'docx']:
                    from docx import Document
                    doc_path = self.file_path.path
                    document = Document(doc_path)
                    content = '\n'.join([para.text for para in document.paragraphs])
                elif self.file_type in ['xlsx', 'xls']:
                    import openpyxl
                    wb = openpyxl.load_workbook(self.file_path.path, read_only=True)
                    content_lines = []
                    for sheet in wb:
                        for row in sheet.iter_rows(values_only=True):
                            content_lines.append('\t'.join(str(cell) if cell is not None else '' for cell in row))
                    content = '\n'.join(content_lines)
                elif self.file_type in ['html', 'htm']:
                    from bs4 import BeautifulSoup
                    with open(self.file_path.path, 'r', encoding='utf-8') as f:
                        soup = BeautifulSoup(f.read(), 'html.parser')
                        content = soup.get_text(separator='\n', strip=True)
                elif self.file_type in ['md', 'txt']:
                    with open(self.file_path.path, 'r', encoding='utf-8') as f:
                        content = f.read()
                elif self.file_type == 'pdf':
                    import PyPDF2
                    with open(self.file_path.path, 'rb') as f:
                        pdf_reader = PyPDF2.PdfReader(f)
                        content = '\n'.join(page.extract_text() for page in pdf_reader.pages)
                else:
                    content = f'不支持的文件类型: {self.file_type}'

                self.file_content = content
                self.save(update_fields=['file_content'])
                return content
            except Exception as e:
                return f'文档解析失败: {e}'
        return ''

    def mark_processed(self):
        """标记为已处理"""
        self.is_processed = True
        self.save(update_fields=['is_processed'])


# 保留原有模型以兼容现有代码（待迁移）
class Doc(models.Model):
    """文档模型（待废弃）"""
    id = models.AutoField(primary_key=True)
    filename = models.CharField(max_length=500, verbose_name='文件名')
    file_content = models.TextField(blank=True, null=True, verbose_name='文件内容')
    file_path = models.FileField(upload_to='prdDocs/', blank=True, null=True, verbose_name='文件路径')
    create_time = models.DateTimeField(auto_now_add=True, verbose_name='创建时间')
    update_time = models.DateTimeField(auto_now=True, verbose_name='更新时间')
    is_case_generated = models.BooleanField(default=False, verbose_name='是否已生成用例')
    project = models.ForeignKey(
        Project,
        on_delete=models.CASCADE,
        related_name='prdDocs',
        verbose_name='所属项目'
    )

    class Meta:
        verbose_name = '文档（旧）'
        verbose_name_plural = '文档（旧）'
        ordering = ['-create_time']

    def __str__(self):
        return f"{self.filename} - {self.project.project_name}"

    @property
    def doc_content(self):
        """获取文档内容"""
        if self.file_content:
            return self.file_content
        elif self.file_path:
            try:
                from docx import Document
                doc_path = self.file_path.path
                document = Document(doc_path)
                # 提取所有段落文本
                content = '\n'.join([para.text for para in document.paragraphs])
                # 将解析的内容保存到file_content字段
                self.file_content = content
                self.save(update_fields=['file_content'])
                return content
            except Exception as e:
                return f'文档解析失败: {e}'
        return ''

    def mark_case_generated(self):
        """标记为已生成用例"""
        self.is_case_generated = True
        self.save(update_fields=['is_case_generated'])


class ApiDoc(models.Model):
    """接口文档模型（待废弃）"""
    id = models.AutoField(primary_key=True)
    filename = models.CharField(max_length=500, verbose_name='文件名')
    version = models.CharField(max_length=50, default='1.0.0', verbose_name='版本号')
    file_content = models.TextField(blank=True, null=True, verbose_name='文件内容')
    file_path = models.FileField(upload_to='apiDocs/', blank=True, null=True, verbose_name='文件路径')
    create_time = models.DateTimeField(auto_now_add=True, verbose_name='创建时间')
    update_time = models.DateTimeField(auto_now=True, verbose_name='更新时间')
    project = models.ForeignKey(
        Project,
        on_delete=models.CASCADE,
        related_name='apiDocs',
        verbose_name='所属项目'
    )

    class Meta:
        verbose_name = '接口文档（旧）'
        verbose_name_plural = '接口文档（旧）'
        ordering = ['-create_time']

    def __str__(self):
        return f"{self.filename} - v{self.version} - {self.project.project_name}"

class ApiInterface(models.Model):
    """API接口模型"""
    METHOD_CHOICES = [
        ('GET', 'GET'),
        ('POST', 'POST'),
        ('PUT', 'PUT'),
        ('DELETE', 'DELETE'),
        ('PATCH', 'PATCH'),
        ('HEAD', 'HEAD'),
        ('OPTIONS', 'OPTIONS'),
    ]

    id = models.AutoField(primary_key=True)
    api_name = models.CharField(max_length=500, verbose_name='接口名称')
    api_path = models.CharField(max_length=1000, blank=True, null=True, verbose_name='接口路径')
    method = models.CharField(max_length=10, choices=METHOD_CHOICES, verbose_name='请求方法')
    request_params = models.TextField(blank=True, null=True, verbose_name='入参')
    response_params = models.TextField(blank=True, null=True, verbose_name='出参')
    remark = models.TextField(blank=True, null=True, verbose_name='备注')
    create_time = models.DateTimeField(auto_now_add=True, verbose_name='创建时间')
    update_time = models.DateTimeField(auto_now=True, verbose_name='更新时间')
    api_doc = models.ForeignKey(
        CommonDoc,
        on_delete=models.CASCADE,
        related_name='api_interfaces',
        verbose_name='所属文档'
    )

    class Meta:
        verbose_name = 'API接口'
        verbose_name_plural = 'API接口'
        ordering = ['-create_time']

    def __str__(self):
        return f"{self.method} {self.api_name}"


class TestData(models.Model):
    """测试数据模型"""
    id = models.AutoField(primary_key=True)
    test_name = models.CharField(max_length=500, verbose_name='测试数据名称')
    test_data_json = models.JSONField(verbose_name='测试数据JSON')
    description = models.TextField(blank=True, null=True, verbose_name='描述')
    is_public = models.BooleanField(default=False, verbose_name='是否公共参数')
    create_time = models.DateTimeField(auto_now_add=True, verbose_name='创建时间')
    update_time = models.DateTimeField(auto_now=True, verbose_name='更新时间')
    api_interface = models.ForeignKey(
        ApiInterface,
        on_delete=models.CASCADE,
        related_name='test_data',
        null=True,
        blank=True,
        verbose_name='所属接口'
    )

    class Meta:
        verbose_name = '测试数据'
        verbose_name_plural = '测试数据'
        ordering = ['-create_time']

    def __str__(self):
        if self.api_interface:
            return f"{self.test_name} - {self.api_interface.api_name}"
        return f"{self.test_name} - 公共参数"
